import cloudinary
import cloudinary.uploader
from PyPDF2 import PdfReader
from docx import Document
import openpyxl
from app.config import settings
from typing import List, Dict
import uuid

cloudinary.config(
    cloud_name=settings.CLOUDINARY_CLOUD_NAME,
    api_key=settings.CLOUDINARY_API_KEY,
    api_secret=settings.CLOUDINARY_API_SECRET
)

class DocumentService:
    @staticmethod
    async def upload_file(file_content: bytes, filename: str, session_id: str) -> Dict:
        """Upload file to Cloudinary and return URL"""
        result = cloudinary.uploader.upload(
            file_content,
            resource_type="raw",
            folder=f"documents/{session_id}",
            public_id=f"{uuid.uuid4()}_{filename}"
        )
        return {
            "url": result["secure_url"],
            "public_id": result["public_id"],
            "filename": filename
        }
    
    @staticmethod
    def extract_text_from_pdf(file_path: str) -> List[Dict[str, any]]:
        """Extract text from PDF with page numbers"""
        reader = PdfReader(file_path)
        pages = []
        for page_num, page in enumerate(reader.pages, start=1):
            text = page.extract_text()
            if text.strip():
                pages.append({
                    "page_number": page_num,
                    "text": text,
                    "total_pages": len(reader.pages)
                })
        return pages
    
    @staticmethod
    def extract_text_from_docx(file_path: str) -> str:
        """Extract text from DOCX"""
        doc = Document(file_path)
        return "\n".join([paragraph.text for paragraph in doc.paragraphs])
    
    @staticmethod
    def extract_text_from_xlsx(file_path: str) -> str:
        """Extract text from Excel"""
        workbook = openpyxl.load_workbook(file_path)
        text_content = []
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                row_text = " ".join([str(cell) for cell in row if cell])
                if row_text.strip():
                    text_content.append(row_text)
        return "\n".join(text_content)
    
    @staticmethod
    def extract_text_from_txt(file_path: str) -> str:
        """Extract text from TXT"""
        with open(file_path, 'r', encoding='utf-8') as f:
            return f.read()
